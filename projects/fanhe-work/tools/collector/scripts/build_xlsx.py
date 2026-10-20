"""把 data/csv30 下的表汇总成一个 Excel 工作簿（每张表一个页签）+ 索引页"""
import csv
import os
import re
import glob
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path

# 工作区根目录由脚本位置推导，盘符随电脑变化时无需修改。
ROOT = Path(__file__).resolve().parent.parent
CSV_ROOT = str(ROOT / 'data' / 'csv30')
OUT = str(ROOT / 'output' / 'qianniu-30d-2026-09-19_2026-10-18.xlsx')

MODULE_ORDER = {'trade-detail': 0, 'trade-refund': 1, 'flow-overview': 2, 'flow-shopsource': 3,
                'flow-itemsource': 4, 'flow-visitor': 5, 'flow-keyword': 6,
                'item-rank': 7, 'item-archives': 8, 'item-macro': 9,
                'content-overview': 10, 'content-live': 11, 'content-channel': 12,
                'qos-core': 13, 'qos-experience': 14, 'qos-performance': 15, 'qos-refund': 16}
MODULE_CN = {
    'trade-detail': '交易', 'trade-refund': '交易-退款',
    'flow-overview': '流量-总览', 'flow-shopsource': '流量-来源', 'flow-itemsource': '流量-商品来源',
    'flow-visitor': '流量-访客', 'flow-keyword': '流量-搜索词',
    'item-rank': '商品-排行', 'item-archives': '商品-360', 'item-macro': '商品-宏观',
    'content-overview': '内容-概况', 'content-live': '内容-直播', 'content-channel': '内容-渠道',
    'qos-core': '服务-核心', 'qos-experience': '服务-体验', 'qos-performance': '服务-绩效', 'qos-refund': '服务-退款',
}


def read_csv(fp):
    with open(fp, encoding='utf-8-sig') as f:
        rows = list(csv.reader(f))
    return rows


def safe_sheet(name, used):
    s = re.sub(r'[\[\]\*\?/\\:]', '_', name)
    s = s[:28]
    base = s
    i = 2
    while s in used:
        s = f'{base[:26]}_{i}'
        i += 1
    used.add(s)
    return s


def main():
    index = {}
    idx_fp = os.path.join(CSV_ROOT, 'INDEX.csv')
    if os.path.exists(idx_fp):
        for r in csv.DictReader(open(idx_fp, encoding='utf-8-sig')):
            index[f"{r['模块']}/{r['文件']}"] = r

    files = sorted(glob.glob(os.path.join(CSV_ROOT, '*', '*.csv')))
    # 只丢掉「没有数据行」的表：判行数，不判列数。
    # 原来写的是 len(r[0]) >= 2，判的是列数，导致只有 1 列的单行表被静默丢弃
    # （qos-core 05/08、trade-refund 02 三张），报表少 3 页且没有任何提示。
    def has_rows(fp):
        rows = read_csv(fp)
        return len(rows) >= 2 and any(any(c.strip() for c in r) for r in rows[1:])
    kept, dropped = [], []
    for p in files:
        (kept if has_rows(p) else dropped).append(p)
    if dropped:
        print('跳过 %d 张无数据行的表：' % len(dropped))
        for p in dropped:
            print('  ' + os.path.relpath(p, CSV_ROOT).replace('\\', '/'))
    files = kept
    files.sort(key=lambda p: (MODULE_ORDER.get(os.path.basename(os.path.dirname(p)), 99), p))

    wb = Workbook()
    wb.remove(wb.active)

    # 索引页
    ws0 = wb.create_sheet('索引')
    head = ['序号', '模块', '表名', '行数', '列数', '来源接口']
    ws0.append(head)
    for i, fp in enumerate(files, 1):
        mod = os.path.basename(os.path.dirname(fp))
        fn = os.path.basename(fp)
        meta = index.get(f'{mod}/{fn}', {})
        rows = read_csv(fp)
        ws0.append([i, MODULE_CN.get(mod, mod), meta.get('中文名', fn),
                    max(0, len(rows) - 1), len(rows[0]) if rows else 0,
                    meta.get('来源接口', '')[:180]])

    hdr_fill = PatternFill('solid', fgColor='1F4E79')
    hdr_font = Font(color='FFFFFF', bold=True, size=10)
    for c in ws0[1]:
        c.fill = hdr_fill
        c.font = hdr_font
    for col, w in zip('ABCDEF', [6, 14, 44, 8, 8, 70]):
        ws0.column_dimensions[col].width = w
    ws0.freeze_panes = 'A2'

    used = {'索引'}
    for i, fp in enumerate(files, 1):
        mod = os.path.basename(os.path.dirname(fp))
        fn = os.path.basename(fp)
        meta = index.get(f'{mod}/{fn}', {})
        label = meta.get('中文名', fn.replace('.csv', ''))
        title = safe_sheet(f'{i:02d}_{label}', used)
        ws = wb.create_sheet(title)
        rows = read_csv(fp)
        if not rows:
            continue
        for r in rows:
            ws.append([str(x)[:200] for x in r])
        for c in ws[1]:
            c.fill = hdr_fill
            c.font = hdr_font
            c.alignment = Alignment(vertical='center')
        ws.freeze_panes = 'A2'
        ncol = len(rows[0])
        for j in range(1, ncol + 1):
            maxlen = max((len(str(rows[k][j - 1])) for k in range(min(len(rows), 60)) if j - 1 < len(rows[k])), default=8)
            ws.column_dimensions[get_column_letter(j)].width = min(46, max(9, maxlen + 2))

    wb.save(OUT)
    print('已生成', OUT)
    print('页签数', len(wb.sheetnames))


if __name__ == '__main__':
    main()
